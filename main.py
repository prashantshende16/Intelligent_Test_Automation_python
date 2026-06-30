import threading
from datetime import datetime
import os
import json
import csv
import io
import openpyxl
from openpyxl import Workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from fastapi import FastAPI, Depends, HTTPException, BackgroundTasks
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func, text, inspect
from typing import List

from database import engine, Base, get_db
import models
import schemas
from agent import run_testing_agent, run_test_execution_agent

# Initialize Database tables
Base.metadata.create_all(bind=engine)

def reset_stuck_tasks():
    from database import SessionLocal
    db = SessionLocal()
    try:
        stuck_tasks = db.query(models.Task).filter(models.Task.status.in_(["crawling", "generating_test_cases", "running_tests", "pending"])).all()
        for task in stuck_tasks:
            task.status = "stopped"
            task.completed_at = datetime.utcnow()
            
        stuck_agents = db.query(models.AgentState).filter(models.AgentState.status.in_({"pending", "running"})).all()
        for state in stuck_agents:
            state.status = "failed"
            state.completed_at = datetime.utcnow()
            
        if stuck_tasks or stuck_agents:
            db.commit()
            print(f"Startup check: Reset {len(stuck_tasks)} stuck tasks and {len(stuck_agents)} active agents to stopped/failed.")
    except Exception as e:
        print(f"Failed to reset stuck tasks on startup: {e}")
    finally:
        db.close()

reset_stuck_tasks()

def ensure_task_auth_columns():
    inspector = inspect(engine)
    if "task_auths" not in inspector.get_table_names():
        return
    columns = {col["name"] for col in inspector.get_columns("task_auths")}
    if "auth_post_login_url" not in columns:
        with engine.begin() as conn:
            conn.execute(text("ALTER TABLE task_auths ADD COLUMN auth_post_login_url VARCHAR"))
    if "auth_next_step" not in columns:
        with engine.begin() as conn:
            conn.execute(text("ALTER TABLE task_auths ADD COLUMN auth_next_step VARCHAR"))
    if "auth_required_fields" not in columns:
        with engine.begin() as conn:
            conn.execute(text("ALTER TABLE task_auths ADD COLUMN auth_required_fields TEXT"))
    if "auth_flow" not in columns:
        with engine.begin() as conn:
            conn.execute(text("ALTER TABLE task_auths ADD COLUMN auth_flow VARCHAR"))

ensure_task_auth_columns()

def ensure_task_columns():
    inspector = inspect(engine)
    if "tasks" in inspector.get_table_names():
        columns = {col["name"] for col in inspector.get_columns("tasks")}
        if "is_mobile" not in columns:
            with engine.begin() as conn:
                conn.execute(text("ALTER TABLE tasks ADD COLUMN is_mobile INTEGER DEFAULT 0"))
        if "ai_model" not in columns:
            with engine.begin() as conn:
                conn.execute(text("ALTER TABLE tasks ADD COLUMN ai_model VARCHAR DEFAULT 'gemini-1.5-flash'"))
        if "user_prompt" not in columns:
            with engine.begin() as conn:
                conn.execute(text("ALTER TABLE tasks ADD COLUMN user_prompt TEXT"))
        if "custom_use_cases_json" not in columns:
            with engine.begin() as conn:
                conn.execute(text("ALTER TABLE tasks ADD COLUMN custom_use_cases_json TEXT"))
        if "page_snapshots_json" not in columns:
            with engine.begin() as conn:
                conn.execute(text("ALTER TABLE tasks ADD COLUMN page_snapshots_json TEXT"))

    if "test_cases" in inspector.get_table_names():
        tc_columns = {col["name"] for col in inspector.get_columns("test_cases")}
        if "page_url" not in tc_columns:
            with engine.begin() as conn:
                conn.execute(text("ALTER TABLE test_cases ADD COLUMN page_url VARCHAR"))
        if "test_type" not in tc_columns:
            with engine.begin() as conn:
                conn.execute(text("ALTER TABLE test_cases ADD COLUMN test_type VARCHAR"))

    if "code_references" in inspector.get_table_names():
        cr_columns = {col["name"] for col in inspector.get_columns("code_references")}
        if "trace_chain_json" not in cr_columns:
            with engine.begin() as conn:
                conn.execute(text("ALTER TABLE code_references ADD COLUMN trace_chain_json TEXT"))

    if "test_errors" in inspector.get_table_names():
        err_columns = {col["name"] for col in inspector.get_columns("test_errors")}
        if "video_path" not in err_columns:
            with engine.begin() as conn:
                conn.execute(text("ALTER TABLE test_errors ADD COLUMN video_path VARCHAR"))

ensure_task_columns()

app = FastAPI(title="AI Website Testing Automation API")

# Enable CORS for frontend app
app.add_middleware(
    CORSMiddleware,
    # FastAPI/Starlette does not allow allow_credentials=True with allow_origins=["*"].
    # Use explicit origins during development.
    allow_origins=[
        "http://localhost:5173",
        "http://127.0.0.1:5173",
        "http://localhost:8000",
        "http://127.0.0.1:8000",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


def get_task_with_counts(task: models.Task, db: Session) -> schemas.TaskResponse:
    """Helper function to calculate counts for a single task."""
    use_case_count = db.query(func.count(models.UseCase.id)).filter(models.UseCase.task_id == task.id).scalar() or 0
    test_case_count = db.query(func.count(models.TestCase.id)).filter(models.TestCase.task_id == task.id).scalar() or 0
    error_count = db.query(func.count(models.TestError.id)).filter(models.TestError.task_id == task.id, models.TestError.severity != "passed", models.TestError.severity != "info").scalar() or 0
    suggestion_count = db.query(func.count(models.Suggestion.id)).filter(models.Suggestion.task_id == task.id).scalar() or 0
    
    return schemas.TaskResponse(
        id=task.id,
        url=task.url,
        status=task.status,
        is_mobile=bool(task.is_mobile),
        ai_model=task.ai_model,
        user_prompt=task.user_prompt,
        custom_use_cases_json=task.custom_use_cases_json,
        page_snapshots_json=task.page_snapshots_json,
        created_at=task.created_at,
        completed_at=task.completed_at,
        use_case_count=use_case_count,
        test_case_count=test_case_count,
        error_count=error_count,
        suggestion_count=suggestion_count
    )

def start_task_thread(task_id: str) -> None:
    thread = threading.Thread(target=run_testing_agent, args=(task_id,))
    thread.daemon = True
    thread.start()

@app.post("/api/tasks", response_model=schemas.TaskResponse)
def create_task(task_in: schemas.TaskCreate, db: Session = Depends(get_db)):
    # Validate codebase path if provided
    if task_in.codebase_path:
        abs_path = os.path.abspath(task_in.codebase_path)
        if not os.path.isdir(abs_path):
            raise HTTPException(status_code=400, detail=f"Codebase directory path '{task_in.codebase_path}' does not exist locally.")
            
    # Create Task record
    db_task = models.Task(
        url=task_in.url,
        status="pending",
        is_mobile=1 if task_in.is_mobile else 0,
        ai_model=task_in.ai_model,
        user_prompt=task_in.user_prompt,
        custom_use_cases_json=task_in.custom_use_cases_json
    )
    db.add(db_task)
    db.commit()
    db.refresh(db_task)
    
    # Save codebase details if provided
    if task_in.codebase_path:
        abs_path = os.path.abspath(task_in.codebase_path)
        db_codebase = models.Codebase(
            task_id=db_task.id,
            local_path=abs_path,
            framework_type="React"
        )
        db.add(db_codebase)

    if getattr(task_in, "auth_required", False):
        db_auth = models.TaskAuth(
            task_id=db_task.id,
            auth_required=1,
            auth_login_url=task_in.auth_login_url,
            auth_post_login_url=task_in.auth_post_login_url,
            auth_username=task_in.auth_username,
            auth_password=task_in.auth_password,
            auth_otp_code=task_in.auth_otp_code,
            auth_otp_hint=task_in.auth_otp_hint,
            auth_flow=task_in.auth_flow,
            auth_next_step=task_in.auth_next_step,
            auth_required_fields=task_in.auth_required_fields,
        )
        db.add(db_auth)

    seed_urls = []
    if getattr(task_in, "seed_urls", None):
        seed_urls = [seed.strip() for seed in task_in.seed_urls if seed and seed.strip()]
    if seed_urls:
        db_seed = models.TaskSeed(
            task_id=db_task.id,
            seed_urls_json=json.dumps(seed_urls)
        )
        db.add(db_seed)
        
    # Pre-populate Agent states for UI tracking
    for agent_name in [
        "Orchestrator", "RouteDiscovery", "HealthCheck", "Login",
        "RolePermission", "UserJourney", "Form", "API",
        "DatabaseIntegrity", "Security", "Accessibility", "Responsive",
        "VisualRegression", "Performance", "CodeCorrelation"
    ]:
        agent_state = models.AgentState(
            task_id=db_task.id,
            agent_name=agent_name,
            status="pending"
        )
        db.add(agent_state)

    # Pre-populate SafetyConfig
    db_safety = models.SafetyConfig(
        task_id=db_task.id,
        protected_usernames_json='["admin","superadmin","root","administrator","sysadmin"]',
        protected_actions_json='["delete","deactivate","change_password","change_role","update","edit","modify","reset_password"]',
        enable_safe_mode=1,
        temp_user_prefix="test_user_",
        cleanup_after_test=1
    )
    db.add(db_safety)
        
    db.commit()
    db.refresh(db_task)
    
    # Spawn background thread to run AI Testing Agent
    start_task_thread(db_task.id)
    
    return get_task_with_counts(db_task, db)

@app.patch("/api/tasks/{task_id}/input", response_model=schemas.TaskResponse)
def update_task_input(task_id: str, task_input: schemas.TaskInputUpdate, db: Session = Depends(get_db)):
    task = db.query(models.Task).filter(models.Task.id == task_id).first()
    if not task:
        raise HTTPException(status_code=404, detail="Task not found")

    auth = db.query(models.TaskAuth).filter(models.TaskAuth.task_id == task_id).first()
    if task_input.auth_required is not None:
        if not auth:
            auth = models.TaskAuth(task_id=task_id, auth_required=int(bool(task_input.auth_required)))
            db.add(auth)
        auth.auth_required = int(bool(task_input.auth_required))
    if auth:
        if task_input.auth_login_url is not None:
            auth.auth_login_url = task_input.auth_login_url
        if task_input.auth_post_login_url is not None:
            auth.auth_post_login_url = task_input.auth_post_login_url
        if task_input.auth_username is not None:
            auth.auth_username = task_input.auth_username
        if task_input.auth_password is not None:
            auth.auth_password = task_input.auth_password
        if task_input.auth_otp_code is not None:
            auth.auth_otp_code = task_input.auth_otp_code
        if task_input.auth_otp_hint is not None:
            auth.auth_otp_hint = task_input.auth_otp_hint
        if task_input.auth_flow is not None:
            auth.auth_flow = task_input.auth_flow
        if task_input.auth_next_step is not None:
            auth.auth_next_step = task_input.auth_next_step
        if task_input.auth_required_fields is not None:
            auth.auth_required_fields = task_input.auth_required_fields

    if task_input.form_values_json is not None:
        form_data = db.query(models.TaskFormData).filter(models.TaskFormData.task_id == task_id).first()
        if not form_data:
            form_data = models.TaskFormData(task_id=task_id, form_values_json=task_input.form_values_json)
            db.add(form_data)
        else:
            form_data.form_values_json = task_input.form_values_json

    if task_input.seed_urls is not None:
        seed_urls = [seed.strip() for seed in task_input.seed_urls if seed and seed.strip()]
        seed_record = db.query(models.TaskSeed).filter(models.TaskSeed.task_id == task_id).first()
        if not seed_record:
            seed_record = models.TaskSeed(task_id=task_id, seed_urls_json=json.dumps(seed_urls))
            db.add(seed_record)
        else:
            seed_record.seed_urls_json = json.dumps(seed_urls)

    if task_input.is_mobile is not None:
        task.is_mobile = 1 if task_input.is_mobile else 0

    if task_input.ai_model is not None:
        task.ai_model = task_input.ai_model

    if task_input.user_prompt is not None:
        task.user_prompt = task_input.user_prompt

    if task_input.custom_use_cases_json is not None:
        task.custom_use_cases_json = task_input.custom_use_cases_json

    task.status = "pending"
    task.completed_at = None
    db.commit()
    db.refresh(task)
    return get_task_with_counts(task, db)

@app.post("/api/tasks/{task_id}/resume", response_model=schemas.TaskResponse)
def resume_task(task_id: str, db: Session = Depends(get_db)):
    task = db.query(models.Task).filter(models.Task.id == task_id).first()
    if not task:
        raise HTTPException(status_code=404, detail="Task not found")
    task.status = "pending"
    task.completed_at = None
    db.commit()
    start_task_thread(task_id)
    return get_task_with_counts(task, db)

@app.post("/api/tasks/{task_id}/stop", response_model=schemas.TaskResponse)
def stop_task(task_id: str, db: Session = Depends(get_db)):
    task = db.query(models.Task).filter(models.Task.id == task_id).first()
    if not task:
        raise HTTPException(status_code=404, detail="Task not found")
    task.status = "stopped"
    task.completed_at = datetime.utcnow()
    
    # Update running agent states to failed/stopped
    agent_states = db.query(models.AgentState).filter(models.AgentState.task_id == task_id).all()
    for state in agent_states:
        if state.status in {"pending", "running"}:
            state.status = "failed"
            state.completed_at = datetime.utcnow()
            
    db.commit()
    db.refresh(task)
    return get_task_with_counts(task, db)


def start_test_execution_thread(task_id: str) -> None:
    thread = threading.Thread(target=run_test_execution_agent, args=(task_id,))
    thread.daemon = True
    thread.start()


@app.post("/api/tasks/{task_id}/start-test", response_model=schemas.TaskResponse)
def start_test_run(task_id: str, db: Session = Depends(get_db)):
    task = db.query(models.Task).filter(models.Task.id == task_id).first()
    if not task:
        raise HTTPException(status_code=404, detail="Task not found")
    
    if task.status in ["crawling", "generating_test_cases", "running_tests"]:
        raise HTTPException(status_code=400, detail=f"Task is in state '{task.status}' and cannot be started.")
        
    task.status = "running_tests"
    task.completed_at = None
    db.commit()
    db.refresh(task)
    
    start_test_execution_thread(task.id)
    return get_task_with_counts(task, db)


@app.post("/api/tasks/{task_id}/stop-test", response_model=schemas.TaskResponse)
def stop_test_run(task_id: str, db: Session = Depends(get_db)):
    return stop_task(task_id, db)

@app.get("/api/tasks", response_model=List[schemas.TaskResponse])
def list_tasks(db: Session = Depends(get_db)):
    tasks = db.query(models.Task).order_by(models.Task.created_at.desc()).all()
    return [get_task_with_counts(task, db) for task in tasks]

@app.get("/api/tasks/{task_id}", response_model=schemas.TaskResponse)
def get_task(task_id: str, db: Session = Depends(get_db)):
    task = db.query(models.Task).filter(models.Task.id == task_id).first()
    if not task:
        raise HTTPException(status_code=404, detail="Task not found")
    return get_task_with_counts(task, db)

@app.get("/api/tasks/{task_id}/details", response_model=schemas.TaskDetailsResponse)
def get_task_details(task_id: str, db: Session = Depends(get_db)):
    # Get task
    task = db.query(models.Task).filter(models.Task.id == task_id).first()
    if not task:
        raise HTTPException(status_code=404, detail="Task not found")
        
    task_resp = get_task_with_counts(task, db)
    
    # Fetch related records
    use_cases = db.query(models.UseCase).filter(models.UseCase.task_id == task_id).order_by(models.UseCase.created_at.asc()).all()
    test_cases = db.query(models.TestCase).filter(models.TestCase.task_id == task_id).order_by(models.TestCase.created_at.asc()).all()
    errors = db.query(models.TestError).filter(models.TestError.task_id == task_id).order_by(models.TestError.created_at.desc()).all()
    suggestions = db.query(models.Suggestion).filter(models.Suggestion.task_id == task_id).order_by(models.Suggestion.created_at.asc()).all()
    
    codebase = db.query(models.Codebase).filter(models.Codebase.task_id == task_id).first()
    auth = db.query(models.TaskAuth).filter(models.TaskAuth.task_id == task_id).first()
    seeds = db.query(models.TaskSeed).filter(models.TaskSeed.task_id == task_id).first()
    agent_states = db.query(models.AgentState).filter(models.AgentState.task_id == task_id).order_by(models.AgentState.started_at.asc()).all()
    auth_state = None
    if auth:
        required_fields = []
        if auth.auth_required_fields:
            try:
                parsed_fields = json.loads(auth.auth_required_fields)
                if isinstance(parsed_fields, list):
                    required_fields = parsed_fields
            except Exception:
                required_fields = [auth.auth_required_fields]
        auth_state = {
            "required": bool(auth.auth_required),
            "flow": auth.auth_flow,
            "next_step": auth.auth_next_step,
            "required_fields": required_fields,
            "login_url": auth.auth_login_url,
            "post_login_url": auth.auth_post_login_url,
            "username": auth.auth_username,
            "password_set": bool((auth.auth_password or "").strip()),
            "otp_set": bool((auth.auth_otp_code or "").strip()),
            "otp_hint": auth.auth_otp_hint,
        }
    safety_config = db.query(models.SafetyConfig).filter(models.SafetyConfig.task_id == task_id).first()
    
    return schemas.TaskDetailsResponse(
        task=task_resp,
        use_cases=use_cases,
        test_cases=test_cases,
        errors=errors,
        suggestions=suggestions,
        codebase=codebase,
        auth=auth,
        auth_state=auth_state,
        seeds=seeds,
        agent_states=agent_states,
        safety_config=safety_config
    )

@app.get("/api/tasks/{task_id}/report.csv")
def download_task_report(task_id: str, db: Session = Depends(get_db)):
    task = db.query(models.Task).filter(models.Task.id == task_id).first()
    if not task:
        raise HTTPException(status_code=404, detail="Task not found")

    use_cases = db.query(models.UseCase).filter(models.UseCase.task_id == task_id).all()
    test_cases = db.query(models.TestCase).filter(models.TestCase.task_id == task_id).all()
    errors = db.query(models.TestError).filter(models.TestError.task_id == task_id).all()
    suggestions = db.query(models.Suggestion).filter(models.Suggestion.task_id == task_id).all()
    codebase = db.query(models.Codebase).filter(models.Codebase.task_id == task_id).first()
    auth = db.query(models.TaskAuth).filter(models.TaskAuth.task_id == task_id).first()
    seeds = db.query(models.TaskSeed).filter(models.TaskSeed.task_id == task_id).first()
    agent_states = db.query(models.AgentState).filter(models.AgentState.task_id == task_id).all()

    output = io.StringIO()
    writer = csv.writer(output)

    def section(title: str):
        writer.writerow([])
        writer.writerow([title.upper()])

    section("Task Summary")
    writer.writerow(["TASK ID", "URL", "STATUS", "CREATED AT", "COMPLETED AT", "TEST CASES", "ERRORS", "SUGGESTIONS"])
    writer.writerow([
        task.id,
        task.url,
        task.status,
        task.created_at.isoformat(),
        task.completed_at.isoformat() if task.completed_at else "",
        len(test_cases),
        len(errors),
        len(suggestions)
    ])

    section("Codebase")
    writer.writerow(["PATH", "FRAMEWORK", "ANALYZED AT"])
    if codebase:
        writer.writerow([
            codebase.local_path,
            codebase.framework_type,
            codebase.analyzed_at.isoformat() if codebase.analyzed_at else ""
        ])
    else:
        writer.writerow(["Not provided", "", ""])

    section("Authentication")
    writer.writerow(["AUTH REQUIRED", "LOGIN URL", "USERNAME", "OTP HINT"])
    if auth:
        writer.writerow([
            str(bool(auth.auth_required)),
            auth.auth_login_url or "",
            auth.auth_username or "",
            auth.auth_otp_hint or ""
        ])
    else:
        writer.writerow(["Not provided", "", "", ""])

    section("Seed URLs")
    writer.writerow(["SEED URL"])
    if seeds and seeds.seed_urls_json:
        try:
            seed_list = json.loads(seeds.seed_urls_json)
        except Exception:
            seed_list = [seeds.seed_urls_json]
        for seed in seed_list:
            writer.writerow([seed])
    else:
        writer.writerow(["Not provided"])

    section("Use Cases")
    writer.writerow(["USE CASE TITLE", "PAGE URL", "DESCRIPTION", "CREATED AT", "USE CASE ID"])
    test_case_page_urls = {}
    for err in errors:
        if err.test_case_id and err.page_url:
            test_case_page_urls[err.test_case_id] = err.page_url
    use_case_page_urls = {}
    for tc in test_cases:
        if not tc.use_case_id:
            continue
        page_url = test_case_page_urls.get(tc.id)
        if page_url:
            use_case_page_urls.setdefault(tc.use_case_id, page_url)
    for err in errors:
        if not err.test_case_id or not err.page_url:
            continue
        tc = next((item for item in test_cases if item.id == err.test_case_id), None)
        if tc and tc.use_case_id:
            use_case_page_urls.setdefault(tc.use_case_id, err.page_url)
    for uc in use_cases:
        writer.writerow([uc.title, use_case_page_urls.get(uc.id, ""), uc.description or "", uc.created_at.isoformat(), uc.id])

    section("Test Cases")
    writer.writerow(["TITLE", "STATUS", "PAGE URL", "EXPECTED RESULT", "ERROR MESSAGE", "STEPS", "EXECUTION TIME", "CREATED AT", "USE CASE ID"])
    for tc in test_cases:
        writer.writerow([
            tc.title,
            tc.status,
            test_case_page_urls.get(tc.id, ""),
            tc.expected_result or "",
            tc.error_message or "",
            tc.steps or "",
            tc.execution_time if tc.execution_time is not None else "",
            tc.created_at.isoformat(),
            tc.use_case_id or "",
        ])

    section("Errors")
    writer.writerow(["MESSAGE", "SEVERITY", "PAGE URL", "SCREENSHOT", "VIDEO", "CREATED AT", "TEST CASE ID"])
    for err in errors:
        writer.writerow([
            err.message,
            err.severity,
            err.page_url,
            err.screenshot_path or "",
            err.video_path or "",
            err.created_at.isoformat(),
            err.test_case_id or "",
        ])

    section("Suggestions")
    writer.writerow(["TITLE", "PRIORITY", "DESCRIPTION", "CREATED AT"])
    for sug in suggestions:
        writer.writerow([sug.title, sug.priority, sug.description or "", sug.created_at.isoformat()])

    section("Agent States")
    writer.writerow(["AGENT", "STATUS", "WARNINGS", "STARTED AT", "COMPLETED AT", "LOG OUTPUT"])
    for agent in agent_states:
        writer.writerow([
            agent.agent_name,
            agent.status,
            str(agent.errors_found),
            agent.started_at.isoformat() if agent.started_at else "",
            agent.completed_at.isoformat() if agent.completed_at else "",
            (agent.log_output or "").replace("\n", " | "),
        ])

    output.seek(0)
    filename = f"{task.id}_test_report.csv"
    headers = {"Content-Disposition": f'attachment; filename="{filename}"'}
    return StreamingResponse(iter([output.getvalue()]), media_type="text/csv", headers=headers)

def generate_task_report_workbook(task_id: str, db: Session, relative_hyperlinks: bool = False) -> Workbook:
    task = db.query(models.Task).filter(models.Task.id == task_id).first()
    if not task:
        raise HTTPException(status_code=404, detail="Task not found")

    use_cases = db.query(models.UseCase).filter(models.UseCase.task_id == task_id).all()
    test_cases = db.query(models.TestCase).filter(models.TestCase.task_id == task_id).all()
    errors = db.query(models.TestError).filter(models.TestError.task_id == task_id).all()
    suggestions = db.query(models.Suggestion).filter(models.Suggestion.task_id == task_id).all()
    codebase = db.query(models.Codebase).filter(models.Codebase.task_id == task_id).first()
    auth = db.query(models.TaskAuth).filter(models.TaskAuth.task_id == task_id).first()
    seeds = db.query(models.TaskSeed).filter(models.TaskSeed.task_id == task_id).first()
    agent_states = db.query(models.AgentState).filter(models.AgentState.task_id == task_id).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Test Report"
    ws.views.sheetView[0].showGridLines = True

    # Style presets
    section_font = Font(name="Segoe UI", size=13, bold=True, color="1F4E79")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    data_font = Font(name="Segoe UI", size=10, color="000000")
    
    thin_border_side = Side(border_style="thin", color="D9D9D9")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    section_border_side = Side(border_style="medium", color="1F4E79")
    section_border = Border(bottom=section_border_side)

    def write_section(title: str):
        if ws.max_row > 1 or (ws.max_row == 1 and ws.cell(row=1, column=1).value is not None):
            ws.append([])
        ws.append([title.upper()])
        row_idx = ws.max_row
        cell = ws.cell(row=row_idx, column=1)
        cell.font = section_font
        cell.border = section_border

    def write_headers(headers: List[str]):
        ws.append(headers)
        row_idx = ws.max_row
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = thin_border
        ws.row_dimensions[row_idx].height = 24

    def write_row(row_data: List):
        processed_row = []
        for val in row_data:
            if isinstance(val, datetime):
                processed_row.append(val.isoformat())
            elif isinstance(val, bool):
                processed_row.append(str(val))
            elif val is None:
                processed_row.append("")
            else:
                processed_row.append(val)

        ws.append(processed_row)
        row_idx = ws.max_row
        for col_idx in range(1, len(processed_row) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[row_idx].height = 20

    def embed_screenshot_in_cell(cell_coordinate: str, screenshot_path: str, row_idx: int):
        if not screenshot_path:
            return
        path_to_try = screenshot_path
        if not os.path.isabs(path_to_try):
            backend_dir = os.path.dirname(os.path.abspath(__file__))
            path_to_try = os.path.join(backend_dir, screenshot_path)
        
        if os.path.exists(path_to_try):
            try:
                from PIL import Image as PILImage
                import io
                
                with PILImage.open(path_to_try) as pil_img:
                    w, h = pil_img.size
                    
                    # Target constraints: max width = 400, max height = 180
                    max_w = 400
                    max_h = 180
                    
                    # Scale to fit within max_w and max_h while preserving aspect ratio
                    ratio = min(max_w / w, max_h / h)
                    target_width = int(w * ratio)
                    target_height = int(h * ratio)
                    
                    try:
                        resampling_filter = PILImage.Resampling.LANCZOS
                    except AttributeError:
                        resampling_filter = PILImage.ANTIALIAS
                        
                    pil_resized = pil_img.resize((target_width, target_height), resampling_filter)
                    
                    img_byte_arr = io.BytesIO()
                    pil_resized.save(img_byte_arr, format='PNG')
                    img_byte_arr.seek(0)
                    
                    img = OpenpyxlImage(img_byte_arr)
                    
                    # Construct URL pointing to backend served screenshot or relative path
                    parts = screenshot_path.replace("\\", "/").split("/")
                    if len(parts) >= 3:
                        filename = parts[-1]
                        task_folder = parts[-2]
                    else:
                        filename = os.path.basename(screenshot_path)
                        task_folder = task_id
                        
                    if relative_hyperlinks:
                        # Inside the ZIP, we place screenshots in the "screenshots" folder next to the Excel file
                        screenshot_url = f"screenshots/{filename}"
                    else:
                        screenshot_url = f"http://localhost:8000/api/screenshots/{task_folder}/{filename}"
                    
                    cell = ws[cell_coordinate]
                    if screenshot_url:
                        cell.value = "Click to view full screenshot"
                        cell.hyperlink = screenshot_url
                        cell.font = Font(name="Segoe UI", size=9, color="0563C1", underline="single", italic=True)
                    else:
                        cell.value = ""
                    
                    # Align text to the bottom center so it is visible below the image
                    cell.alignment = Alignment(horizontal="center", vertical="bottom")
                    
                    ws.add_image(img, cell_coordinate)
                    ws.row_dimensions[row_idx].height = 160  # Plenty of room for 180px image + text
            except Exception as e:
                ws[cell_coordinate].value = f"Error: {str(e)}"

    # Section 1: Task Summary
    write_section("Task Summary")
    write_headers(["TASK ID", "URL", "STATUS", "CREATED AT", "COMPLETED AT", "TEST CASES", "ERRORS", "SUGGESTIONS"])
    write_row([
        task.id,
        task.url,
        task.status,
        task.created_at,
        task.completed_at,
        len(test_cases),
        len(errors),
        len(suggestions)
    ])

    # Section 2: Codebase
    write_section("Codebase")
    write_headers(["PATH", "FRAMEWORK", "ANALYZED AT"])
    if codebase:
        write_row([
            codebase.local_path,
            codebase.framework_type,
            codebase.analyzed_at
        ])
    else:
        write_row(["Not provided", "", ""])

    # Section 3: Authentication
    write_section("Authentication")
    write_headers(["AUTH REQUIRED", "LOGIN URL", "USERNAME", "OTP HINT"])
    if auth:
        write_row([
            bool(auth.auth_required),
            auth.auth_login_url or "",
            auth.auth_username or "",
            auth.auth_otp_hint or ""
        ])
    else:
        write_row(["Not provided", "", "", ""])

    # Section 4: Seed URLs
    write_section("Seed URLs")
    write_headers(["SEED URL"])
    if seeds and seeds.seed_urls_json:
        try:
            seed_list = json.loads(seeds.seed_urls_json)
        except Exception:
            seed_list = [seeds.seed_urls_json]
        for seed in seed_list:
            write_row([seed])
    else:
        write_row(["Not provided"])

    # Section 5: Use Cases
    write_section("Use Cases")
    write_headers(["USE CASE TITLE", "PAGE URL", "DESCRIPTION", "CREATED AT", "USE CASE ID"])
    
    test_case_page_urls = {}
    for err in errors:
        if err.test_case_id and err.page_url:
            test_case_page_urls[err.test_case_id] = err.page_url
            
    use_case_page_urls = {}
    for tc in test_cases:
        if not tc.use_case_id:
            continue
        page_url = test_case_page_urls.get(tc.id)
        if page_url:
            use_case_page_urls.setdefault(tc.use_case_id, page_url)
            
    for err in errors:
        if not err.test_case_id or not err.page_url:
            continue
        tc = next((item for item in test_cases if item.id == err.test_case_id), None)
        if tc and tc.use_case_id:
            use_case_page_urls.setdefault(tc.use_case_id, err.page_url)
            
    for uc in use_cases:
        write_row([
            uc.title,
            use_case_page_urls.get(uc.id, ""),
            uc.description or "",
            uc.created_at,
            uc.id
        ])

    # Section 6: Test Cases
    write_section("Test Cases")
    write_headers(["TITLE", "STATUS", "PAGE URL", "SCREENSHOT", "EXPECTED RESULT", "ERROR MESSAGE", "STEPS", "EXECUTION TIME", "CREATED AT", "USE CASE ID"])
    
    tc_screenshots = {}
    for err in errors:
        if err.test_case_id and err.screenshot_path:
            tc_screenshots[err.test_case_id] = err.screenshot_path

    for tc in test_cases:
        screenshot_val = tc_screenshots.get(tc.id, "")
        write_row([
            tc.title,
            tc.status,
            test_case_page_urls.get(tc.id, ""),
            screenshot_val,
            tc.expected_result or "",
            tc.error_message or "",
            tc.steps or "",
            tc.execution_time if tc.execution_time is not None else "",
            tc.created_at,
            tc.use_case_id or "",
        ])
        if screenshot_val:
            row_idx = ws.max_row
            embed_screenshot_in_cell(f"D{row_idx}", screenshot_val, row_idx)

    # Section 7: Errors
    write_section("Errors")
    write_headers(["MESSAGE", "SEVERITY", "PAGE URL", "SCREENSHOT", "VIDEO", "CREATED AT", "TEST CASE ID"])
    for err in errors:
        screenshot_val = err.screenshot_path or ""
        video_val = err.video_path or ""
        video_url = ""
        if video_val:
            filename = os.path.basename(video_val)
            if relative_hyperlinks:
                video_url = f"videos/{filename}"
            else:
                video_url = f"http://localhost:8000/api/videos/{task_id}/{filename}"
        write_row([
            err.message,
            err.severity,
            err.page_url,
            screenshot_val,
            video_url,
            err.created_at,
            err.test_case_id or "",
        ])
        row_idx = ws.max_row
        if screenshot_val:
            embed_screenshot_in_cell(f"D{row_idx}", screenshot_val, row_idx)
        if video_url:
            cell = ws[f"E{row_idx}"]
            cell.value = "Click to view video recording"
            cell.hyperlink = video_url
            cell.font = Font(name="Segoe UI", size=10, color="0563C1", underline="single", italic=True)

    # Section 8: Suggestions
    write_section("Suggestions")
    write_headers(["TITLE", "PRIORITY", "DESCRIPTION", "CREATED AT"])
    for sug in suggestions:
        write_row([
            sug.title,
            sug.priority,
            sug.description or "",
            sug.created_at
        ])

    # Section 9: Agent States
    write_section("Agent States")
    write_headers(["AGENT", "STATUS", "WARNINGS", "STARTED AT", "COMPLETED AT", "LOG OUTPUT"])
    for agent in agent_states:
        write_row([
            agent.agent_name,
            agent.status,
            str(agent.errors_found),
            agent.started_at,
            agent.completed_at,
            (agent.log_output or "").replace("\n", " | "),
        ])

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > 100:
                val_str = val_str[:100]
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 50)

    # Explicitly set Column D (SCREENSHOT) to a wider size so images fit nicely
    ws.column_dimensions["D"].width = 55
    return wb

@app.get("/api/tasks/{task_id}/report.xlsx")
def download_task_report_xlsx(task_id: str, db: Session = Depends(get_db)):
    task = db.query(models.Task).filter(models.Task.id == task_id).first()
    if not task:
        raise HTTPException(status_code=404, detail="Task not found")

    wb = generate_task_report_workbook(task_id, db, relative_hyperlinks=False)
    file_stream = io.BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)

    filename = f"{task.id}_test_report.xlsx"
    headers = {"Content-Disposition": f'attachment; filename="{filename}"'}
    return StreamingResponse(
        file_stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers
    )

@app.get("/api/tasks/{task_id}/report.zip")
def download_task_report_zip(task_id: str, db: Session = Depends(get_db)):
    import zipfile
    task = db.query(models.Task).filter(models.Task.id == task_id).first()
    if not task:
        raise HTTPException(status_code=404, detail="Task not found")

    # 1. Generate the Excel file with relative links
    wb = generate_task_report_workbook(task_id, db, relative_hyperlinks=True)
    excel_stream = io.BytesIO()
    wb.save(excel_stream)
    excel_stream.seek(0)

    # 2. Query errors to find screenshots
    errors = db.query(models.TestError).filter(models.TestError.task_id == task_id).all()

    # 3. Create the ZIP in memory
    zip_stream = io.BytesIO()
    with zipfile.ZipFile(zip_stream, "w", zipfile.ZIP_DEFLATED) as zip_file:
        # Write Excel file to ZIP root
        excel_filename = f"{task.id}_test_report.xlsx"
        zip_file.writestr(excel_filename, excel_stream.getvalue())

        # Write each screenshot to the screenshots/ folder inside the ZIP
        added_screenshots = set()
        backend_dir = os.path.dirname(os.path.abspath(__file__))

        for err in errors:
            if err.screenshot_path:
                path_to_try = err.screenshot_path
                if not os.path.isabs(path_to_try):
                    path_to_try = os.path.join(backend_dir, err.screenshot_path)

                if os.path.exists(path_to_try) and os.path.isfile(path_to_try):
                    filename = os.path.basename(path_to_try)
                    if filename not in added_screenshots:
                        # Write to "screenshots/{filename}" in ZIP
                        zip_file.write(path_to_try, arcname=f"screenshots/{filename}")
                        added_screenshots.add(filename)

        # Write each video to the videos/ folder inside the ZIP
        added_videos = set()
        for err in errors:
            if err.video_path:
                path_to_try = err.video_path
                if not os.path.isabs(path_to_try):
                    path_to_try = os.path.join(backend_dir, err.video_path)

                if os.path.exists(path_to_try) and os.path.isfile(path_to_try):
                    filename = os.path.basename(path_to_try)
                    if filename not in added_videos:
                        # Write to "videos/{filename}" in ZIP
                        zip_file.write(path_to_try, arcname=f"videos/{filename}")
                        added_videos.add(filename)

    zip_stream.seek(0)

    zip_filename = f"{task.id}_test_report.zip"
    headers = {"Content-Disposition": f'attachment; filename="{zip_filename}"'}
    return StreamingResponse(
        zip_stream,
        media_type="application/zip",
        headers=headers
    )

@app.delete("/api/tasks/{task_id}")
def delete_task(task_id: str, db: Session = Depends(get_db)):
    task = db.query(models.Task).filter(models.Task.id == task_id).first()
    if not task:
        raise HTTPException(status_code=404, detail="Task not found")
        
    db.delete(task)
    db.commit()
    return {"message": f"Task {task_id} successfully deleted"}

@app.get("/api/screenshots/{task_id}/{filename}")
def get_screenshot(task_id: str, filename: str):
    safe_filename = os.path.basename(filename)
    candidate_dirs = [
        os.path.abspath(os.path.join("screenshots", task_id)),
        os.path.abspath(os.path.join(os.path.dirname(__file__), "screenshots", task_id)),
    ]

    for screenshot_dir in candidate_dirs:
        screenshot_path = os.path.abspath(os.path.join(screenshot_dir, safe_filename))
        if screenshot_path.startswith(screenshot_dir) and os.path.isfile(screenshot_path):
            return FileResponse(screenshot_path, media_type="image/png")

    raise HTTPException(status_code=404, detail="Screenshot not found")

@app.get("/api/videos/{task_id}/{filename}")
def get_video(task_id: str, filename: str):
    safe_filename = os.path.basename(filename)
    candidate_dirs = [
        os.path.abspath(os.path.join("videos", task_id)),
        os.path.abspath(os.path.join(os.path.dirname(__file__), "videos", task_id)),
    ]

    for video_dir in candidate_dirs:
        video_path = os.path.abspath(os.path.join(video_dir, safe_filename))
        if video_path.startswith(video_dir) and os.path.isfile(video_path):
            return FileResponse(video_path, media_type="video/webm")

    raise HTTPException(status_code=404, detail="Video not found")

@app.get("/api/dashboard/stats", response_model=schemas.DashboardStats)
def get_dashboard_stats(db: Session = Depends(get_db)):
    # Task Counts
    total_tasks = db.query(func.count(models.Task.id)).scalar() or 0
    total_test_cases = db.query(func.count(models.TestCase.id)).scalar() or 0
    total_errors = db.query(func.count(models.TestError.id)).filter(models.TestError.severity != "passed", models.TestError.severity != "info").scalar() or 0
    total_suggestions = db.query(func.count(models.Suggestion.id)).scalar() or 0
    
    # Calculate Success Rate
    passed_tests = db.query(func.count(models.TestCase.id)).filter(models.TestCase.status == "passed").scalar() or 0
    failed_tests = db.query(func.count(models.TestCase.id)).filter(models.TestCase.status == "failed").scalar() or 0
    total_run_tests = passed_tests + failed_tests
    
    success_rate = 100.0
    if total_run_tests > 0:
        success_rate = round((passed_tests / total_run_tests) * 100, 1)
        
    # Group tasks by status
    status_query = db.query(models.Task.status, func.count(models.Task.id)).group_by(models.Task.status).all()
    status_counts = {status: count for status, count in status_query}
    
    # Initialize missing statuses to zero for client predictability
    for status in ["pending", "crawling", "generating_test_cases", "planned", "running_tests", "completed", "failed", "stopped"]:
        if status not in status_counts:
            status_counts[status] = 0
            
    return schemas.DashboardStats(
        total_tasks=total_tasks,
        total_test_cases=total_test_cases,
        total_errors=total_errors,
        total_suggestions=total_suggestions,
        success_rate=success_rate,
        status_counts=status_counts
    )
